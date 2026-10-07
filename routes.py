import os
import uuid
from datetime import datetime
from pathlib import Path

from flask import (
    Blueprint,
    render_template,
    request,
    redirect,
    url_for,
    flash,
    session,
)

from werkzeug.utils import secure_filename
from openpyxl import load_workbook

from auth import admin_required
from database import db
from models import School, Student, Admin, Loan
from forms import ExcelImportForm, StudentForm, SchoolForm
from excel_validator import validate_workbook


admin = Blueprint("admin", __name__, url_prefix="/admin")


IMPORT_DIR = (
    Path(__file__).resolve().parent
    / "instance"
    / "imports"
)

ALLOWED_EXTENSIONS = {".xlsx"}

MAX_EXCEL_SIZE = 10 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 100
PREVIEW_ROWS = 15


def allowed_excel(filename):
    return (
        Path(filename).suffix.lower()
        in ALLOWED_EXTENSIONS
    )


def safe_import_filename(original_name):
    suffix = Path(original_name).suffix.lower()

    random_name = (
        f"{uuid.uuid4().hex}{suffix}"
    )

    return secure_filename(random_name)


def inspect_worksheet(ws):
    """
    بررسی یک Sheet بدون وارد کردن اطلاعات به دیتابیس.
    """

    row_count = 0
    max_columns = 0
    preview_rows = []

    for row in ws.iter_rows(values_only=True):

        values = list(row)

        # ردیف کاملاً خالی را نادیده می‌گیریم
        if not any(
            value is not None
            and str(value).strip() != ""
            for value in values
        ):
            continue

        row_count += 1

        if len(values) > MAX_COLUMNS:
            raise ValueError(
                f"تعداد ستون‌های Sheet «{ws.title}» "
                f"بیشتر از حد مجاز ({MAX_COLUMNS}) است."
            )

        max_columns = max(
            max_columns,
            len(values)
        )

        if len(preview_rows) < PREVIEW_ROWS:
            preview_rows.append(values)

        if row_count > MAX_ROWS:
            raise ValueError(
                f"تعداد ردیف‌های Sheet «{ws.title}» "
                f"بیشتر از حد مجاز ({MAX_ROWS}) است."
            )

    headers = []

    if preview_rows:
        headers = [
            "" if value is None else str(value)
            for value in preview_rows[0]
        ]

    return {
        "name": ws.title,
        "row_count": row_count,
        "column_count": max_columns,
        "headers": headers,
        "rows": preview_rows,
        "empty": row_count == 0,
    }


@admin.route("/")
@admin_required
def dashboard():

    now = datetime.utcnow()

    school_count = db.session.scalar(
        db.select(
            db.func.count(School.id)
        )
    ) or 0

    student_count = db.session.scalar(
        db.select(
            db.func.count(Student.id)
        )
    ) or 0

    admin_count = db.session.scalar(
        db.select(
            db.func.count(Admin.id)
        )
    ) or 0

    active_loan_count = db.session.scalar(
        db.select(
            db.func.count(Loan.id)
        ).where(
            Loan.status == "active"
        )
    ) or 0

    overdue_loan_count = db.session.scalar(
        db.select(
            db.func.count(Loan.id)
        ).where(
            Loan.status == "active",
            Loan.due_at < now
        )
    ) or 0

    returned_loan_count = db.session.scalar(
        db.select(
            db.func.count(Loan.id)
        ).where(
            Loan.status == "returned"
        )
    ) or 0

    recent_loans = db.session.scalars(
        db.select(Loan)
        .order_by(
            Loan.created_at.desc()
        )
        .limit(6)
    ).all()

    return render_template(
        "dashboard.html",
        school_count=school_count,
        student_count=student_count,
        admin_count=admin_count,
        active_loan_count=active_loan_count,
        overdue_loan_count=overdue_loan_count,
        returned_loan_count=returned_loan_count,
        recent_loans=recent_loans,
    )


@admin.route("/schools")
@admin_required
def schools():
    schools_list = db.session.scalars(
        db.select(School).order_by(School.name)
    ).all()

    return render_template(
        "admin/schools.html",
        schools=schools_list
    )


@admin.route("/schools/new", methods=["GET", "POST"])
@admin_required
def school_create():

    form = SchoolForm()

    if form.validate_on_submit():

        name = form.name.data.strip()
        code = form.code.data.strip()

        existing = db.session.execute(
            db.select(School).where(
                School.code == code
            )
        ).scalar_one_or_none()

        if existing:
            flash(
                "این کد مدرسه قبلاً ثبت شده است.",
                "error"
            )

            return render_template(
                "admin/school_create.html",
                form=form
            )

        school = School(
            name=name,
            code=code
        )

        db.session.add(school)
        db.session.commit()

        flash(
            "مدرسه با موفقیت ثبت شد.",
            "success"
        )

        return redirect(
            url_for("admin.schools")
        )

    return render_template(
        "admin/school_create.html",
        form=form
    )


@admin.route("/students")
@admin_required
def students():
    query = request.args.get("q", "").strip()

    statement = (
        db.select(Student)
        .join(School, Student.school_id == School.id)
        .order_by(
            Student.last_name,
            Student.first_name
        )
    )

    if query:
        pattern = f"%{query}%"

        statement = statement.where(
            db.or_(
                Student.first_name.ilike(pattern),
                Student.last_name.ilike(pattern),
                Student.student_code.ilike(pattern),
                School.name.ilike(pattern)
            )
        )

    students_list = db.session.scalars(
        statement
    ).all()

    student_count = db.session.scalar(
        db.select(db.func.count(Student.id))
    ) or 0

    active_loan_count = db.session.scalar(
        db.select(db.func.count(Loan.id)).where(
            Loan.status == "active"
        )
    ) or 0

    return render_template(
        "admin/students.html",
        students=students_list,
        student_count=student_count,
        active_loan_count=active_loan_count,
        query=query
    )


@admin.route("/students/<int:student_id>")
@admin_required
def student_detail(student_id):

    student = db.session.get(
        Student,
        student_id
    )

    if student is None:
        flash(
            "عضو موردنظر پیدا نشد.",
            "error"
        )

        return redirect(
            url_for("admin.students")
        )

    loan_count = db.session.scalar(
        db.select(
            db.func.count(Loan.id)
        ).where(
            Loan.student_id == student.id
        )
    ) or 0

    active_loan_count = db.session.scalar(
        db.select(
            db.func.count(Loan.id)
        ).where(
            Loan.student_id == student.id,
            Loan.status == "active"
        )
    ) or 0

    loan_history = db.session.scalars(
        db.select(Loan)
        .where(
            Loan.student_id == student.id
        )
        .order_by(
            Loan.created_at.desc()
        )
    ).all()

    return render_template(
        "admin/student_detail.html",
        student=student,
        loan_count=loan_count,
        active_loan_count=active_loan_count,
        loan_history=loan_history
    )


@admin.route("/students/new", methods=["GET", "POST"])
@admin_required
def student_create():

    schools = db.session.scalars(
        db.select(School).order_by(School.name)
    ).all()

    form = StudentForm()

    form.school_id.choices = [
        (school.id, school.name)
        for school in schools
    ]

    if not schools:
        flash(
            "ابتدا باید حداقل یک مدرسه ثبت کنید.",
            "error"
        )

    if form.validate_on_submit():

        student_code = form.student_code.data.strip()

        existing = db.session.execute(
            db.select(Student).where(
                Student.student_code == student_code
            )
        ).scalar_one_or_none()

        if existing:
            flash(
                "این کد دانش‌آموزی قبلاً ثبت شده است.",
                "error"
            )

            return render_template(
                "admin/student_create.html",
                form=form
            )

        selected_school = db.session.get(
            School,
            form.school_id.data
        )

        if selected_school is None:
            flash(
                "مدرسه انتخاب‌شده معتبر نیست.",
                "error"
            )

            return render_template(
                "admin/student_create.html",
                form=form
            )

        student = Student(
            first_name=form.first_name.data.strip(),
            last_name=form.last_name.data.strip(),
            student_code=student_code,
            grade=form.grade.data.strip(),
            class_name=form.class_name.data.strip(),
            school_id=selected_school.id
        )

        db.session.add(student)
        db.session.commit()

        flash(
            "عضو با موفقیت ثبت شد.",
            "success"
        )

        return redirect(
            url_for("admin.students")
        )

    return render_template(
        "admin/student_create.html",
        form=form
    )


@admin.route("/loans")
@admin_required
def loans():
    return render_template("admin/loans.html")


@admin.route("/loans/new", methods=["GET", "POST"])
@admin_required
def loan_create():
    from forms import LoanForm

    form = LoanForm()

    if form.validate_on_submit():
        flash(
            "فرم با موفقیت بررسی شد. ثبت نهایی پس از اتصال کتاب‌ها و اعضای واقعی فعال می‌شود.",
            "success"
        )

        return redirect(url_for("admin.loans"))

    return render_template(
        "admin/loan_create.html",
        form=form
    )


@admin.route("/books")
@admin_required
def books():
    return render_template("admin/books.html")


@admin.route("/books/import/cancel", methods=["POST"])
@admin_required
def books_import_cancel():
    stored_name = session.pop("excel_import_file", None)
    session.pop("excel_import_original_name", None)

    if stored_name:
        import_path = IMPORT_DIR / stored_name

        try:
            if import_path.exists() and import_path.is_file():
                import_path.unlink()
        except OSError:
            pass

    flash(
        "فرآیند ورود فایل لغو شد و فایل موقت حذف شد.",
        "success"
    )

    return redirect(
        url_for("admin.books_import")
    )


@admin.route(
    "/books/import",
    methods=["GET", "POST"]
)
@admin_required
def books_import():

    form = ExcelImportForm()

    if request.method == "GET":
        return render_template(
            "admin/books_import.html",
            form=form
        )

    if not form.validate_on_submit():

        flash(
            "درخواست نامعتبر است. لطفاً دوباره تلاش کنید.",
            "error"
        )

        return redirect(
            url_for("admin.books_import")
        )

    uploaded = request.files.get(
        "excel_file"
    )

    if not uploaded:

        flash(
            "لطفاً یک فایل Excel انتخاب کنید.",
            "error"
        )

        return redirect(
            url_for("admin.books_import")
        )

    original_name = (
        uploaded.filename or ""
    ).strip()

    if not original_name:

        flash(
            "نام فایل معتبر نیست.",
            "error"
        )

        return redirect(
            url_for("admin.books_import")
        )

    if not allowed_excel(original_name):

        flash(
            "فعلاً فقط فایل‌های Excel با پسوند .xlsx قابل قبول هستند.",
            "error"
        )

        return redirect(
            url_for("admin.books_import")
        )

    try:

        uploaded.stream.seek(
            0,
            os.SEEK_END
        )

        file_size = uploaded.stream.tell()

        uploaded.stream.seek(0)

    except Exception:

        flash(
            "امکان بررسی اندازه فایل وجود ندارد.",
            "error"
        )

        return redirect(
            url_for("admin.books_import")
        )

    if file_size <= 0:

        flash(
            "فایل خالی است.",
            "error"
        )

        return redirect(
            url_for("admin.books_import")
        )

    if file_size > MAX_EXCEL_SIZE:

        flash(
            "حجم فایل بیشتر از ۱۰ مگابایت است.",
            "error"
        )

        return redirect(
            url_for("admin.books_import")
        )

    IMPORT_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    stored_name = safe_import_filename(
        original_name
    )

    stored_path = (
        IMPORT_DIR / stored_name
    )

    workbook = None

    try:

        uploaded.save(
            stored_path
        )

        workbook = load_workbook(
            filename=stored_path,
            read_only=True,
            data_only=True
        )

        sheet_names = workbook.sheetnames

        if not sheet_names:

            raise ValueError(
                "فایل هیچ Sheet قابل استفاده‌ای ندارد."
            )

        validation = validate_workbook(
            workbook
        )

        sheets = []

        for sheet_name in sheet_names:

            ws = workbook[sheet_name]

            sheet_info = inspect_worksheet(
                ws
            )

            sheets.append(
                sheet_info
            )

        session["excel_import_file"] = (
            stored_name
        )

        session["excel_import_original_name"] = (
            original_name
        )

        return render_template(
            "admin/books_import_preview.html",
            original_name=original_name,
            file_size=file_size,
            sheets=sheets,
            validation=validation
        )

    except Exception as exc:

        flash(
            f"خواندن فایل Excel انجام نشد: {exc}",
            "error"
        )

        return redirect(
            url_for("admin.books_import")
        )

    finally:

        if workbook is not None:

            try:
                workbook.close()
            except Exception:
                pass

        if (
            "stored_path" in locals()
            and stored_path.exists()
            and "excel_import_file" not in session
        ):

            try:
                stored_path.unlink()
            except OSError:
                pass
