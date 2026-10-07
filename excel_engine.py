from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook


MAX_ROWS = 5000
MAX_COLUMNS = 100
MAX_HEADER_LENGTH = 200
PREVIEW_ROWS = 15


def normalize_header(value: Any) -> str:
    """
    تبدیل Header به متن استاندارد برای تحلیل.
    هنوز هیچ نگاشتی به فیلدهای Book انجام نمی‌شود.
    """
    if value is None:
        return ""

    text = str(value)

    text = text.replace("\u200c", " ")
    text = " ".join(text.split())

    return text.strip()


def normalize_cell(value: Any) -> str:
    """
    تبدیل مقدار سلول به متن قابل نمایش و مقایسه.
    """
    if value is None:
        return ""

    return str(value).strip()


def is_non_empty_row(values: list[Any]) -> bool:
    return any(
        normalize_cell(value) != ""
        for value in values
    )


def analyze_worksheet(ws) -> dict:
    """
    تحلیل ساختاری یک Worksheet.

    این تابع عمداً هیچ فرضی درباره ستون‌های کتاب ندارد.
    """

    raw_rows = []
    non_empty_rows = 0
    empty_rows = 0
    max_columns_seen = 0

    for row in ws.iter_rows(values_only=True):

        values = list(row)

        max_columns_seen = max(
            max_columns_seen,
            len(values)
        )

        if len(values) > MAX_COLUMNS:
            raise ValueError(
                f"تعداد ستون‌های Sheet «{ws.title}» "
                f"بیشتر از حد مجاز ({MAX_COLUMNS}) است."
            )

        if not is_non_empty_row(values):
            empty_rows += 1
            continue

        non_empty_rows += 1

        if non_empty_rows > MAX_ROWS:
            raise ValueError(
                f"تعداد ردیف‌های Sheet «{ws.title}» "
                f"بیشتر از حد مجاز ({MAX_ROWS}) است."
            )

        if len(raw_rows) < PREVIEW_ROWS:
            raw_rows.append(values)

    if not raw_rows:
        return {
            "name": ws.title,
            "empty": True,
            "row_count": 0,
            "column_count": 0,
            "header_count": 0,
            "duplicate_headers": [],
            "empty_headers": [],
            "headers": [],
            "rows": [],
            "empty_rows": empty_rows,
            "status": "empty",
        }

    header_values = raw_rows[0]

    headers = [
        normalize_header(value)
        for value in header_values
    ]

    empty_headers = [
        index + 1
        for index, header in enumerate(headers)
        if not header
    ]

    seen = {}
    duplicate_headers = []

    for index, header in enumerate(headers, start=1):

        if not header:
            continue

        if len(header) > MAX_HEADER_LENGTH:
            raise ValueError(
                f"Header ستون {index} در Sheet "
                f"«{ws.title}» بیش از {MAX_HEADER_LENGTH} "
                f"کاراکتر است."
            )

        key = header.casefold()

        if key in seen:
            duplicate_headers.append({
                "name": header,
                "first_column": seen[key],
                "duplicate_column": index,
            })
        else:
            seen[key] = index

    return {
        "name": ws.title,
        "empty": False,
        "row_count": non_empty_rows,
        "column_count": max_columns_seen,
        "header_count": len(headers),
        "duplicate_headers": duplicate_headers,
        "empty_headers": empty_headers,
        "headers": headers,
        "rows": raw_rows,
        "empty_rows": empty_rows,
        "status": (
            "warning"
            if duplicate_headers or empty_headers
            else "ready"
        ),
    }


def analyze_excel(file_path: str | Path) -> dict:
    """
    تحلیل کل فایل Excel.

    فقط تحلیل انجام می‌شود و هیچ تغییری در دیتابیس ایجاد نمی‌شود.
    """

    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(
            "فایل Excel پیدا نشد."
        )

    if not file_path.is_file():
        raise ValueError(
            "مسیر واردشده یک فایل نیست."
        )

    workbook = None

    try:

        workbook = load_workbook(
            filename=file_path,
            read_only=True,
            data_only=True
        )

        sheet_names = workbook.sheetnames

        if not sheet_names:
            raise ValueError(
                "فایل هیچ Sheet قابل استفاده‌ای ندارد."
            )

        sheets = []

        for sheet_name in sheet_names:

            worksheet = workbook[sheet_name]

            sheets.append(
                analyze_worksheet(worksheet)
            )

        total_rows = sum(
            sheet["row_count"]
            for sheet in sheets
        )

        warning_count = sum(
            1
            for sheet in sheets
            if sheet["status"] == "warning"
        )

        empty_count = sum(
            1
            for sheet in sheets
            if sheet["status"] == "empty"
        )

        return {
            "filename": file_path.name,
            "sheet_count": len(sheets),
            "total_rows": total_rows,
            "warning_count": warning_count,
            "empty_count": empty_count,
            "sheets": sheets,
        }

    finally:

        if workbook is not None:

            try:
                workbook.close()
            except Exception:
                pass
